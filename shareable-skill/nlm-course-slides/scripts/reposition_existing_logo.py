#!/usr/bin/env python3
"""
Inspect or reposition the locally-added logo in existing PPTX files in place.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Cm


OLD_LEFT = Cm(40.99)
OLD_TOP = Cm(0.30)
NEW_LEFT = Cm(40.26)
NEW_TOP = Cm(23.36)
LOGO_WIDTH = Cm(4.10)
LOGO_HEIGHT = Cm(1.19)
POSITION_TOLERANCE = Cm(0.3)
SIZE_TOLERANCE = Cm(0.05)


@dataclass
class ShapeInfo:
    slide_index: int
    shape_index: int
    name: str
    left: int
    top: int
    width: int
    height: int


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("target", help="PPTX file or directory containing PPTX files")
    parser.add_argument("--inspect", action="store_true", help="Print picture-shape positions instead of editing")
    parser.add_argument(
        "--pattern",
        default="*水印版.pptx",
        help="Glob pattern used when target is a directory",
    )
    return parser


def _within(value: int, expected: int, tolerance: int) -> bool:
    return abs(int(value) - int(expected)) <= int(tolerance)


def _is_logo_candidate(shape) -> bool:
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return False
    return (
        _within(shape.width, LOGO_WIDTH, SIZE_TOLERANCE)
        and _within(shape.height, LOGO_HEIGHT, SIZE_TOLERANCE)
        and (
            (
                _within(shape.left, OLD_LEFT, POSITION_TOLERANCE)
                and _within(shape.top, OLD_TOP, POSITION_TOLERANCE)
            )
            or (
                _within(shape.left, NEW_LEFT, POSITION_TOLERANCE)
                and _within(shape.top, NEW_TOP, POSITION_TOLERANCE)
            )
        )
    )


def iter_pptx_files(target: Path, pattern: str) -> list[Path]:
    if target.is_file():
        return [target]
    return sorted(path for path in target.glob(pattern) if path.is_file())


def inspect_file(path: Path) -> dict[str, object]:
    prs = Presentation(str(path))
    pictures: list[ShapeInfo] = []
    candidates: list[ShapeInfo] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            info = ShapeInfo(
                slide_index=slide_index,
                shape_index=shape_index,
                name=shape.name,
                left=int(shape.left),
                top=int(shape.top),
                width=int(shape.width),
                height=int(shape.height),
            )
            pictures.append(info)
            if _is_logo_candidate(shape):
                candidates.append(info)
    return {
        "file": str(path),
        "picture_count": len(pictures),
        "candidate_count": len(candidates),
        "pictures": [asdict(item) for item in pictures],
        "candidates": [asdict(item) for item in candidates],
    }


def reposition_file(path: Path) -> dict[str, object]:
    prs = Presentation(str(path))
    updated = 0
    matched = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not _is_logo_candidate(shape):
                continue
            matched += 1
            if int(shape.left) != int(NEW_LEFT) or int(shape.top) != int(NEW_TOP):
                shape.left = NEW_LEFT
                shape.top = NEW_TOP
                updated += 1
    if matched == 0:
        raise RuntimeError(f"No logo candidate found in {path}")
    prs.save(str(path))
    return {
        "file": str(path),
        "matched": matched,
        "updated": updated,
    }


def main() -> int:
    args = build_parser().parse_args()
    target = Path(args.target).resolve()
    files = iter_pptx_files(target, args.pattern)
    if not files:
        raise SystemExit(f"No PPTX files found for target: {target}")
    if args.inspect:
        payload = [inspect_file(path) for path in files]
    else:
        payload = [reposition_file(path) for path in files]
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
